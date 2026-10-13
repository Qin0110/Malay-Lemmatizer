#!/usr/bin/env python
"""Build a labelling sheet from a sentence pool, and read it back.

    python crawl/make_goldsheet.py                  data/gold_pool.jsonl -> data/gold_sheet.xlsx
    python crawl/make_goldsheet.py --collect        sheet -> data/gold.jsonl

    python crawl/make_goldsheet.py --pool data/val_pool.jsonl \\
        --sheet data/val_sheet.xlsx --out data/val.jsonl [--collect]

Every token is pre-filled by the rules; only fix corrected_lemma where the
proposal is wrong. Save and close Excel before --collect.
"""

import argparse
import json
import sys
import time
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import rules                                                      # noqa: E402

GOLD_POOL = ROOT / "data" / "gold_pool.jsonl"
GOLD_SHEET = ROOT / "data" / "gold_sheet.xlsx"
GOLD_JSONL = ROOT / "data" / "gold.jsonl"

CHECK = ("oov", "dict_check", "ambiguous", "deriv")
WARN = ("affix", "redup", "elong", "name")


def build():
    if not GOLD_POOL.exists():
        sys.exit("No %s." % GOLD_POOL.relative_to(ROOT))
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        sys.exit("openpyxl is needed: pip install openpyxl")

    print("loading MALINDO...")
    rows = [json.loads(l) for l in GOLD_POOL.open(encoding="utf-8") if l.strip()]

    prior = prior_corrections()
    if prior:
        archive = GOLD_SHEET.with_name(
            "%s_%s.xlsx" % (GOLD_SHEET.stem, time.strftime("%Y%m%d_%H%M%S")))
        GOLD_SHEET.rename(archive)
        print("Carrying forward %d corrections (old sheet kept as %s)\n"
              % (len(prior), archive.name))

    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(["sent_id", "sentence", "token", "proposed_lemma",
               "corrected_lemma", "rule"])

    counts = defaultdict(int)
    for row in rows:
        for i, e in enumerate(rules.label(row["text"])):
            counts[e["src"]] += 1
            ws.append([row["id"], row["text"] if i == 0 else "", e["t"], e["l"],
                       prior.get((row["id"], e["t"].lower())), e["src"]])

    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, w in {"A": 11, "B": 60, "C": 16, "D": 18, "E": 18, "F": 11}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    hot = PatternFill("solid", fgColor="FCE4D6")
    warn = PatternFill("solid", fgColor="FFF2CC")
    for r in range(2, ws.max_row + 1):
        rule = ws.cell(row=r, column=6).value
        fill = hot if rule in CHECK else warn if rule in WARN else None
        if fill:
            for c in range(1, 7):
                ws.cell(row=r, column=c).fill = fill
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
    wb.save(GOLD_SHEET)

    n = sum(counts.values())
    print("%d sentences, %d tokens -> %s\n" % (len(rows), n, GOLD_SHEET.relative_to(ROOT)))
    for rule, c in sorted(counts.items(), key=lambda x: -x[1]):
        flag = "  <- check these" if rule in CHECK else ""
        print("  %-11s %6d %6.1f%%%s" % (rule, c, 100 * c / n, flag))
    print("\nFill corrected_lemma only where proposed_lemma is wrong, then run "
          "the same command with --collect.")


def prior_corrections():
    if not GOLD_SHEET.exists():
        return {}
    try:
        from openpyxl import load_workbook
        ws = load_workbook(GOLD_SHEET, data_only=True).active
    except Exception:
        return {}
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and len(r) > 4 and r[0] and r[2] and r[4] and str(r[4]).strip():
            out[(r[0], str(r[2]).lower())] = str(r[4]).strip()
    return out


def cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def collect():
    if not GOLD_SHEET.exists():
        sys.exit("No %s yet." % GOLD_SHEET.relative_to(ROOT))
    from openpyxl import load_workbook

    pool = {}
    for line in GOLD_POOL.open(encoding="utf-8"):
        if line.strip():
            row = json.loads(line)
            pool[row["id"]] = row
    lex = rules.malindo()[0]

    ws = load_workbook(GOLD_SHEET, data_only=True).active
    per_sent, edits = defaultdict(list), 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        sent_id, _, token, proposed, corrected, rule = (cell(v) for v in r[:6])
        if corrected:
            edits += 1
        lemma = corrected or proposed or token
        per_sent[sent_id].append({
            "t": token,
            "l": lemma if rule == "name" and not corrected else lemma.lower(),
            "src": "human" if corrected else rule,
        })

    out = []
    for sent_id, lemmas in per_sent.items():
        row = pool.get(sent_id)
        if not row:
            continue
        out.append({
            "id": sent_id,
            "text": row["text"],
            "tokens": [e["t"] for e in lemmas],
            "candidates": {e["t"]: lex.get(e["t"].lower(), []) for e in lemmas},
            "lemmas": lemmas,
            "lemma_text": " ".join(e["l"] for e in lemmas),
            "band": row.get("band"),
        })

    with GOLD_JSONL.open("w", encoding="utf-8") as f:
        for row in out:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    n_tok = sum(len(r["lemmas"]) for r in out)
    print("%d sentences, %d tokens -> %s"
          % (len(out), n_tok, GOLD_JSONL.relative_to(ROOT)))
    print("%d tokens corrected by hand (%.1f%%)" % (edits, 100 * edits / max(n_tok, 1)))
    if edits == 0:
        print("No corrections at all - was the sheet saved and closed?")


if __name__ == "__main__":
    p = argparse.ArgumentParser(description="Build or collect a labelling sheet.")
    p.add_argument("--collect", action="store_true",
                   help="read the finished sheet back into the --out file")
    p.add_argument("--pool", metavar="JSONL", help="default data/gold_pool.jsonl")
    p.add_argument("--sheet", metavar="XLSX", help="default data/gold_sheet.xlsx")
    p.add_argument("--out", metavar="JSONL", help="default data/gold.jsonl")
    a = p.parse_args()

    given = [a.pool, a.sheet, a.out]
    if any(given) and not all(given):
        p.error("--pool, --sheet and --out must be given together")
    if a.pool:
        def at(x):
            x = Path(x)
            return x if x.is_absolute() else (ROOT / x).resolve()
        GOLD_POOL, GOLD_SHEET, GOLD_JSONL = at(a.pool), at(a.sheet), at(a.out)
        if not GOLD_POOL.exists():
            sys.exit("No %s. Run crawl/make_valpool.py first." % GOLD_POOL)

    collect() if a.collect else build()
