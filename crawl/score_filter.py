import sys
from collections import Counter
from pathlib import Path

DATA = Path(__file__).resolve().parent.parent / "data"
CHECK_FILE = DATA / "filter_check.tsv"
CHECK_XLSX = DATA / "filter_check.xlsx"


def read_xlsx():
    """Columns: filter_said | hits | is_really_rojak_Y_or_N | text"""
    from openpyxl import load_workbook
    ws = load_workbook(CHECK_XLSX, data_only=True).active
    rows, blank, bad = [], 0, []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        verdict = str(row[0]).strip().upper()
        label = ("" if row[2] is None else str(row[2])).strip().upper()
        if not label:
            blank += 1
        elif label in ("Y", "N"):
            rows.append((verdict, label))
        else:
            bad.append((n, row[2]))
    return rows, blank, bad


def read_tsv():
    rows, blank, bad = [], 0, []
    with CHECK_FILE.open(encoding="utf-8-sig") as f:
        header = next(f, "")
        if "is_really_rojak" not in header:
            sys.exit("Unexpected header. Regenerate with --validate.")
        for n, line in enumerate(f, start=2):
            parts = line.rstrip("\n").rstrip("\r").split("\t")
            if len(parts) < 4:
                continue
            verdict, label = parts[0].strip().upper(), parts[1].strip().upper()
            if not label:
                blank += 1
            elif label in ("Y", "N"):
                rows.append((verdict, label))
            else:
                bad.append((n, parts[1]))
    return rows, blank, bad


def main():
    # Prefer the xlsx - that is the one people actually label.
    if CHECK_XLSX.exists():
        rows, blank, bad = read_xlsx()
        print("reading %s\n" % CHECK_XLSX.name)
    elif CHECK_FILE.exists():
        rows, blank, bad = read_tsv()
        print("reading %s\n" % CHECK_FILE.name)
    else:
        sys.exit("No filter_check file. Run: "
                 "python crawl/build_corpus.py --validate")

    if bad:
        print("Rows with something other than Y/N (ignored):")
        for n, val in bad[:10]:
            print("  line %d: %r" % (n, val))
        print()

    if blank:
        print("%d rows still blank - fill every row for a valid figure.\n" % blank)

    if not rows:
        sys.exit("Nothing scored yet - fill the Y/N column first.")

    c = Counter(rows)
    kept_y = c[("KEPT", "Y")]
    kept_n = c[("KEPT", "N")]
    rej_y = c[("REJECTED", "Y")]
    rej_n = c[("REJECTED", "N")]

    print("%-26s %5s %5s" % ("", "rojak", "not"))
    print("-" * 38)
    print("%-26s %5d %5d" % ("filter KEPT", kept_y, kept_n))
    print("%-26s %5d %5d" % ("filter REJECTED", rej_y, rej_n))
    print("-" * 38)

    if kept_y + kept_n:
        precision = kept_y / (kept_y + kept_n)
        print("precision  %.1f%%   (%d/%d kept rows are really rojak)"
              % (100 * precision, kept_y, kept_y + kept_n))
    if kept_y + rej_y:
        recall = kept_y / (kept_y + rej_y)
        print("recall     %.1f%%   (%d/%d rojak rows were kept)"
              % (100 * recall, kept_y, kept_y + rej_y))

    print()
    if kept_y + kept_n and kept_y / (kept_y + kept_n) < 0.75:
        print("Precision below 75%: the corpus is contaminated. Raise the bar:")
        print("    python crawl/build_corpus.py --min-hits 3 --validate")
    elif kept_y + rej_y and kept_y / (kept_y + rej_y) < 0.60:
        print("Recall below 60%: too much good data is being discarded. Try:")
        print("    python crawl/build_corpus.py --min-hits 1 --validate")
    else:
        print("Filter looks usable. Note both numbers for the report, then")
        print("move on to labelling data/gold_pool.jsonl.")

    print("\nNOTE for the report: the REJECTED rows here are a sample of what")
    print("the filter threw out, so recall is measured only against sentences")
    print("that survived cleaning and the length filter - not against every")
    print("comment crawled. Say so when you quote the number.")


if __name__ == "__main__":
    main()
